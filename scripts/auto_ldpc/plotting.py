"""
Plotting utilities for LDPC automation.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from scripts.auto_ldpc.state import StateStore


def _import_matplotlib():
    try:
        import matplotlib.pyplot as plt  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("未安装 matplotlib, 无法绘图。请运行 `pip install matplotlib`.") from exc
    return plt


def _load_baseline(
    baseline_cfg: Dict[str, Optional[str]],
    workspace_root: Path,
) -> Optional[Tuple[List[float], List[float]]]:
    path_value = baseline_cfg.get("path")
    if not path_value:
        return None
    baseline_path = Path(path_value)
    if not baseline_path.is_absolute():
        baseline_path = workspace_root / baseline_path
    if not baseline_path.exists():
        return None

    sheet = baseline_cfg.get("sheet") or "baseline"
    snr_col = baseline_cfg.get("snr_column") or "SNR"
    fer_col = baseline_cfg.get("fer_column") or "FER"

    try:
        from openpyxl import load_workbook  # type: igno
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("未安装 openpyxl，无法读取基准数据。请运行 `pip install openpyxl`.") from exc

    wb = load_workbook(baseline_path, data_only=True)
    if sheet not in wb.sheetnames:
        return None
    ws = wb[sheet]
    headers = {}
    snr_values: List[float] = []
    fer_values: List[float] = []
    for row in ws.iter_rows(values_only=True):
        if not headers:
            for idx, value in enumerate(row):
                if isinstance(value, str):
                    headers[value.strip()] = idx
            continue
        if not headers:
            continue
        snr_idx = headers.get(snr_col)
        fer_idx = headers.get(fer_col)
        if snr_idx is None or fer_idx is None:
            continue
        snr = row[snr_idx]
        fer = row[fer_idx]
        if snr is None or fer is None:
            continue
        try:
            snr_values.append(float(snr))
            fer_values.append(float(fer))
        except (TypeError, ValueError):
            continue
    if not snr_values:
        return None
    paired = sorted(zip(snr_values, fer_values), key=lambda x: x[0])
    snr_sorted, fer_sorted = zip(*paired)
    return list(snr_sorted), list(fer_sorted)



def _group_runs_by_matrix(
    runs: Sequence[Dict[str, object]]
) -> Dict[str, List[Tuple[float, float]]]:
    grouped: Dict[str, List[Tuple[float, float]]] = {}
    for run in runs:
        matrix_id = str(run["matrix_id"])
        fer = run.get("fer")
        snr = run.get("snr")
        if fer is None or snr is None:
            continue
        try:
            fer_value = float(fer)
            snr_value = float(snr)
        except (TypeError, ValueError):
            continue
        grouped.setdefault(matrix_id, []).append((snr_value, fer_value))
    for matrix_id in grouped:
        grouped[matrix_id].sort(key=lambda item: item[0])
    return grouped


def plot_top_results(
    state: StateStore,
    matrix_ids: Sequence[str],
    output_path: Path,
    workspace_root: Path,
    baseline_cfg: Optional[Dict[str, Optional[str]]] = None,
) -> Optional[Path]:
    if not matrix_ids:
        return None
    runs = state.fetch_awgn_runs_for_matrices(matrix_ids)
    if not runs:
        return None
    data = _group_runs_by_matrix(runs)
    if not data:
        return None

    plt = _import_matplotlib()
    baseline = None
    if baseline_cfg:
        try:
            baseline = _load_baseline(baseline_cfg, workspace_root)
        except RuntimeError as exc:
            print(f"[WARN] {exc}")
            baseline = None

    plt.figure(figsize=(8, 5))
    for matrix_id in matrix_ids:
        points = data.get(matrix_id)
        if not points:
            continue
        snrs = [p[0] for p in points]
        fers = [p[1] for p in points]
        plt.semilogy(snrs, fers, marker="o", label=f"Matrix {matrix_id}")

    if baseline and baseline[0]:
        plt.semilogy(
            baseline[0],
            baseline[1],
            linestyle="--",
            marker="s",
            color="black",
            label="Baseline",
        )

    plt.xlabel("SNR")
    plt.ylabel("FER")
    plt.title("Top Matrices SNR-FER Curves")
    plt.grid(True, which="both", linestyle="--", linewidth=0.5)
    plt.legend()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    plt.tight_layout()
    plt.savefig(output_path)
    plt.close()
    return output_path
